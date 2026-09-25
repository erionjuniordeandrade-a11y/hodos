# -*- coding: utf-8 -*-
"""Shared table layout, read/write and validation for the resident-authored lesson kit
(Teach-back program unit 1D, docs/plans/2026-09-25-teach-back-program.md §3).

The DOCX has one fixed, label/value table per section. Every label is a unique
Portuguese string, so reading is just: walk every table row in the document and
match the left cell's text against this module's field list. Building writes the
same tables with empty right-hand cells for the resident to fill in Word.

No terminal knowledge is assumed of the resident; this module is only ever run by
an agent or the site owner, never by the resident themselves.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Optional

from docx import Document
from docx.shared import Pt

N_STEPS = 5

INSTRUCTIONS = (
    "Este arquivo é um modelo para você escrever uma aula curta de anatomia para o Hodos. "
    "Você não precisa saber nada de terminal, de programação ou de Git: preencha os espaços "
    "das tabelas abaixo, diretamente aqui no Word, e devolva o arquivo salvo (.docx) para quem "
    "pediu a aula. Não precisa mexer nos títulos em negrito das tabelas, só no texto ao lado deles. "
    "A aula tem cinco “Passos”; cada um é um trecho de explicação (Texto) mais uma frase "
    "descrevendo o que a pessoa deveria estar vendo no visualizador 3D naquele momento (Cena). Em "
    "“Fontes”, escreva uma referência por linha, usando um PMID (só números, por exemplo "
    "33615216, o número do PubMed) ou um DOI (por exemplo 10.1093/brain/awt163). Se não souber o "
    "PMID ou o DOI de um artigo, procure o título no PubMed (pubmed.ncbi.nlm.nih.gov) e copie o "
    "número que aparece na página do artigo."
)

CREDIT_NOTE = (
    "Atenção sobre autoria: o seu nome completo, escrito no campo “Autor”, vai aparecer "
    "como crédito de autoria na aula quando ela for revisada e publicada no Hodos. A aula permanece "
    "marcada como rascunho (“draft”) até a revisão do responsável; ela não fica visível aos "
    "residentes antes disso."
)


def _passo_section(n: int):
    return (
        f"Passo {n}",
        [
            (f"step{n}_title", f"Passo {n} – Título"),
            (f"step{n}_text", f"Passo {n} – Texto"),
            (f"step{n}_scene", f"Passo {n} – Cena (o que mostrar)"),
            (f"step{n}_sources", f"Passo {n} – Fontes (PMID ou DOI, uma por linha)"),
        ],
    )


# (section_heading, [(field_key, label_text), ...]) — the fixed table structure.
SECTIONS = [
    ("Identificação", [
        ("title", "Título da aula"),
        ("author", "Autor (nome completo)"),
        ("date", "Data"),
    ]),
    ("Pergunta de abertura", [("opening_question", "Pergunta de abertura")]),
    ("Resumo", [("summary", "Resumo")]),
    ("Objetivos", [
        ("goal_1", "Objetivo 1"),
        ("goal_2", "Objetivo 2"),
        ("goal_3", "Objetivo 3"),
    ]),
    *[_passo_section(n) for n in range(1, N_STEPS + 1)],
    ("Pergunta final e resposta", [
        ("final_question", "Pergunta final"),
        ("final_answer", "Resposta"),
    ]),
    ("Referências", [("references", "Referências (lista completa, uma por linha)")]),
]

LABEL_BY_FIELD = {field: label for _, rows in SECTIONS for field, label in rows}
FIELD_BY_LABEL = {label: field for field, label in LABEL_BY_FIELD.items()}
ALL_FIELDS = list(LABEL_BY_FIELD.keys())


def build_template() -> Document:
    doc = Document()
    doc.add_heading("Modelo de aula — Hodos (rascunho para revisão)", level=0)

    p = doc.add_paragraph(INSTRUCTIONS)
    p.paragraph_format.space_after = Pt(12)

    note = doc.add_paragraph()
    note_run = note.add_run(CREDIT_NOTE)
    note_run.bold = True
    note.paragraph_format.space_after = Pt(16)

    for heading, rows in SECTIONS:
        doc.add_heading(heading, level=1)
        table = doc.add_table(rows=len(rows), cols=2)
        table.style = "Table Grid"
        table.autofit = True
        for i, (_field, label) in enumerate(rows):
            label_cell, value_cell = table.cell(i, 0), table.cell(i, 1)
            label_run = label_cell.paragraphs[0].add_run(label)
            label_run.bold = True
            value_cell.text = ""
        doc.add_paragraph()
    return doc


_EXTRA_ROW_RE = re.compile(r"^Passo\s+\d+", re.IGNORECASE)


def read_fields(doc: Document) -> dict:
    """Flat field_key -> text (embedded newlines preserved for multi-line cells).

    Rows the template does not define but that look like a step ("Passo 6 – ...")
    are listed under "_unknown_rows", and labels found more than once under
    "_duplicate_rows", so validate() can report them instead of silently dropping
    or overwriting what the resident wrote."""
    values: dict = {}
    unknown: list = []
    duplicates: list = []
    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) < 2:
                continue
            label = row.cells[0].text.strip()
            field = FIELD_BY_LABEL.get(label)
            if field is None:
                if _EXTRA_ROW_RE.match(label) and label not in unknown:
                    unknown.append(label)
                continue
            if field in values and label not in duplicates:
                duplicates.append(label)
            values[field] = row.cells[1].text.strip()
    values["_unknown_rows"] = unknown
    values["_duplicate_rows"] = duplicates
    return values


def write_fields(doc: Document, data: dict) -> None:
    """Fill an already-built template's value cells from a field_key -> text|list dict."""
    label_to_text = {}
    for field, value in data.items():
        label = LABEL_BY_FIELD.get(field)
        if label is None:
            continue
        text = "\n".join(value) if isinstance(value, (list, tuple)) else str(value)
        label_to_text[label] = text
    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) < 2:
                continue
            label = row.cells[0].text.strip()
            if label in label_to_text:
                row.cells[1].text = label_to_text[label]


_PMID_RE = re.compile(r"^\d{3,9}$")
# A DOI never ends in sentence punctuation pasted from prose ("...awt163.").
_DOI_RE = re.compile(r"^10\.\d{4,9}/\S*[^\s.,;:]$", re.IGNORECASE)
_DOI_URL_RE = re.compile(r"^https?://(dx\.)?doi\.org/(10\.\d{4,9}/\S*[^\s.,;:])$", re.IGNORECASE)


@dataclass
class Citation:
    kind: str  # 'pmid' | 'doi'
    value: str
    raw: str


def classify_citation(line: str) -> Optional[Citation]:
    line = line.strip()
    if not line:
        return None
    if _PMID_RE.match(line):
        return Citation("pmid", line, line)
    doi_url = _DOI_URL_RE.match(line)
    if doi_url:
        return Citation("doi", doi_url.group(2), line)
    if line.lower().startswith("doi:"):
        candidate = line[4:].strip()
        if _DOI_RE.match(candidate):
            return Citation("doi", candidate, line)
        return None
    if _DOI_RE.match(line):
        return Citation("doi", line, line)
    return None


def validate(fields: dict) -> list:
    """Returns a list of Portuguese error messages; empty list means the DOCX is valid."""
    errors: list = []

    def require(field: str, message: str):
        if not fields.get(field, "").strip():
            errors.append(message)

    for label in fields.get("_unknown_rows", []):
        errors.append(
            f"A linha “{label}” não faz parte do modelo (a aula tem {N_STEPS} Passos). "
            "Junte esse conteúdo a um dos Passos existentes e apague a linha extra."
        )
    for label in fields.get("_duplicate_rows", []):
        errors.append(
            f"A linha “{label}” aparece mais de uma vez no documento. Deixe só uma e apague a cópia."
        )
    require("title", "Preencha o campo “Título da aula”.")
    author = fields.get("author", "").strip()
    if not author:
        errors.append("Preencha o campo “Autor (nome completo)”.")
    elif len(author.split()) < 2:
        errors.append(
            "O campo “Autor (nome completo)” precisa do nome completo (nome e sobrenome), não só um nome."
        )
    require("date", "Preencha o campo “Data”.")
    require("opening_question", "Preencha a “Pergunta de abertura”.")
    require("summary", "Preencha o “Resumo”.")
    for n in (1, 2, 3):
        require(f"goal_{n}", f"Preencha o “Objetivo {n}”.")

    for n in range(1, N_STEPS + 1):
        title_field, text_field, scene_field, sources_field = (
            f"step{n}_title", f"step{n}_text", f"step{n}_scene", f"step{n}_sources",
        )
        missing = [
            label for field, label in (
                (title_field, "Título"), (text_field, "Texto"), (scene_field, "Cena"),
            ) if not fields.get(field, "").strip()
        ]
        if missing:
            errors.append(
                f"O Passo {n} está incompleto: preencha {', '.join(missing)} antes de enviar."
            )
        sources_raw = fields.get(sources_field, "").strip()
        if not sources_raw:
            errors.append(
                f"O Passo {n} precisa de pelo menos uma fonte (PMID ou DOI) em “Fontes”."
            )
        else:
            for line in sources_raw.splitlines():
                line = line.strip()
                if not line:
                    continue
                if classify_citation(line) is None:
                    errors.append(
                        f"No Passo {n}, a linha “{line}” em “Fontes” não parece um PMID (só "
                        "números) nem um DOI (começa com 10. e tem uma barra). Corrija ou remova essa linha."
                    )

    require("final_question", "Preencha a “Pergunta final”.")
    require("final_answer", "Preencha a “Resposta” (da pergunta final).")
    require("references", "Preencha as “Referências” (lista completa).")

    return errors
